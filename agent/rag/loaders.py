"""Parse uploads (.pdf, .docx, .txt, .md) into pages of plain text."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

SUPPORTED = (".pdf", ".docx", ".txt", ".md")


@dataclass(frozen=True)
class Page:
    number: int
    text: str


class UnsupportedFileError(ValueError):
    pass


def load_pages(path: Path) -> list[Page]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [Page(i + 1, p.extract_text() or "") for i, p in enumerate(reader.pages)]
    if suffix == ".docx":
        import docx

        d = docx.Document(str(path))
        lines = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return [Page(1, "\n".join(lines))]
    if suffix in (".txt", ".md"):
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "utf-16", "cp1251"):
            try:
                return [Page(1, raw.decode(enc))]
            except UnicodeDecodeError:
                continue
        return [Page(1, raw.decode("utf-8", errors="replace"))]
    raise UnsupportedFileError(f"unsupported file type {suffix!r}; use {', '.join(SUPPORTED)}")


def detect_lang(text: str) -> str:
    """'ka' when Georgian letters dominate, else 'en'."""
    ka = sum(1 for ch in text if "ა" <= ch <= "ჿ")
    latin = sum(1 for ch in text if ch.isascii() and ch.isalpha())
    return "ka" if ka > latin else "en"
